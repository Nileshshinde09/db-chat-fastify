import asyncio
import csv
from pathlib import Path
from openpyxl import Workbook


def _safe(v):
    """Prevent CSV/Excel formula injection."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


async def write_csv(
    columns: list[str],
    rows: list[list],
    path: Path,
    summary: str = "",
) -> tuple[int, int]:
    count = 0
    with path.open("w", newline="", encoding="utf-8") as f:
        if summary:
            for line in summary.splitlines():
                f.write(f"# {line}\n")
            f.write("\n")
        w = csv.writer(f)
        w.writerow(columns)
        for row in rows:
            w.writerow([_safe(c) for c in row])
            count += 1
            if count % 1000 == 0:
                await asyncio.sleep(0)
    return count, path.stat().st_size


def write_xlsx(
    columns: list[str],
    rows: list[list],
    path: Path,
    summary: str = "",
) -> tuple[int, int]:
    wb = Workbook(write_only=True)

    if summary:
        ws_sum = wb.create_sheet("Summary")
        ws_sum.append([summary])

    ws = wb.create_sheet("Report")
    ws.append(columns)
    count = 0
    for row in rows:
        ws.append([_safe(c) for c in row])
        count += 1
    wb.save(path)
    return count, path.stat().st_size